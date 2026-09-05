#!/usr/bin/env python3
"""
verify_workbooks.py - confirm every collection workbook matches the register.

Run it from anywhere:

    python verify_workbooks.py

WHY THIS EXISTS.  On 14 August 2026 the DS-division register moved from 331 rows
to 340 (the Survey Department's 2025 revision) and the codes changed from the
generated CEN-001 style to the official KA1 style.  A regeneration of the 243
workbooks was interrupted part-way, leaving a MIXTURE: some workbooks on the new
340-row layout, some still on the old one.

That mixture is the dangerous state, because nothing about it looks wrong.  A
workbook with 331 rows opens fine, sorts fine and prints fine - it is simply
missing nine divisions and keyed on codes the database no longer knows.  Data
entered into it would import as unmatched rows, or worse, silently key to
nothing.

So this script does not ask "did the generator print success".  It opens every
workbook and checks what is actually in it.

WHAT IT CHECKS, per workbook, on both period tabs:
  1. the row count equals that province's row count in dsd_register.csv
  2. every ds_code in column A exists in the register
  3. no ds_code is a retired or legacy code (CEN-001 style, or the three
     retired codes EAS-008 / CEN-032 / CEN-034)
  4. the division names match the register for those codes
  5. the two period tabs agree with each other

EXIT CODE is 0 if everything passes and 1 if anything fails, so you can also use
it in a script.  Nothing is modified - this only reads.
"""

import csv
import os
import sys
from collections import defaultdict

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl is not installed.  Run:  pip install openpyxl")

# Resolve paths relative to this file if it sits in design/templates,
# otherwise fall back to the current directory.
HERE = os.path.dirname(os.path.abspath(__file__))
CANDIDATES = [
    HERE,
    os.path.join(HERE, "design", "templates"),
    os.getcwd(),
    os.path.join(os.getcwd(), "design", "templates"),
]
TEMPLATES = next(
    (c for c in CANDIDATES if os.path.isdir(os.path.join(c, "generated"))), None
)
if TEMPLATES is None:
    sys.exit(
        "Could not find the 'generated' folder.\n"
        "Put this file in design/templates/ and run it from there."
    )
GENERATED = os.path.join(TEMPLATES, "generated")

REG_CANDIDATES = [
    os.path.join(TEMPLATES, "..", "ingestion", "dsd_register.csv"),
    os.path.join(HERE, "..", "ingestion", "dsd_register.csv"),
    os.path.join(HERE, "dsd_register.csv"),
    os.path.join(os.getcwd(), "dsd_register.csv"),
    os.path.join(os.getcwd(), "design", "ingestion", "dsd_register.csv"),
]
REGISTER = next(
    (os.path.abspath(c) for c in REG_CANDIDATES if os.path.isfile(c)), None
)
if REGISTER is None:
    sys.exit(
        "Could not find dsd_register.csv.  Looked in:\n  "
        + "\n  ".join(os.path.abspath(c) for c in REG_CANDIDATES)
    )

# Retired codes must never appear.  Each was a division split in two; a retired
# code resolving to anything would mean roughly half the area it used to mean.
RETIRED = {"EAS-008", "CEN-032", "CEN-034"}

# ---------------------------------------------------------------- the register
by_province = defaultdict(dict)          # province -> {code: name}
all_codes = {}                           # code -> name
with open(REGISTER, encoding="utf-8-sig") as fh:
    for row in csv.DictReader(fh):
        by_province[row["province"]][row["ds_code"]] = row["ds_division"]
        all_codes[row["ds_code"]] = row["ds_division"]

print(f"register: {REGISTER}")
print(f"          {len(all_codes)} divisions across {len(by_province)} provinces")
for p in sorted(by_province):
    print(f"            {p:16s} {len(by_province[p])}")
print()

# Folder name on disk uses underscores: "North_Central" -> "North Central"
def province_of(folder):
    name = folder.replace("_", " ")
    return name if name in by_province else None


# ----------------------------------------------------------------- the check
problems = []      # (workbook, message)
checked = 0

for folder in sorted(os.listdir(GENERATED)):
    fdir = os.path.join(GENERATED, folder)
    if not os.path.isdir(fdir):
        continue
    province = province_of(folder)
    if province is None:
        problems.append((folder, f"folder does not match any province in the register"))
        continue

    expected = by_province[province]

    for fname in sorted(os.listdir(fdir)):
        if not fname.endswith(".xlsx") or fname.startswith("~$"):
            continue
        path = os.path.join(fdir, fname)
        label = f"{folder}/{fname}"
        checked += 1

        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:                      # noqa: BLE001
            problems.append((label, f"could not open: {exc}"))
            continue

        period_tabs = [s for s in wb.sheetnames if "-" in s and s[0].isdigit()]
        if not period_tabs:
            problems.append((label, "no period tabs found"))
            wb.close()
            continue

        per_tab_codes = {}
        for tab in period_tabs:
            ws = wb[tab]
            codes = []
            # Data starts at row 5; row 4 is the locked EXAMPLE row.
            for r in ws.iter_rows(min_row=5, min_col=1, max_col=2, values_only=True):
                code = r[0]
                if code is None or str(code).strip() == "":
                    continue
                codes.append((str(code).strip(), str(r[1]).strip() if r[1] else ""))
            per_tab_codes[tab] = codes

            n = len(codes)
            if n != len(expected):
                problems.append(
                    (label, f"tab '{tab}' has {n} rows, register says {len(expected)}")
                )

            seen = set()
            for code, name in codes:
                if code in RETIRED:
                    problems.append((label, f"tab '{tab}' uses RETIRED code {code}"))
                elif code not in expected:
                    hint = " (looks like an old generated code)" if "-" in code else ""
                    problems.append(
                        (label, f"tab '{tab}' has unknown code {code}{hint}")
                    )
                elif name and name != expected[code]:
                    problems.append(
                        (label,
                         f"tab '{tab}' code {code} is named '{name}', "
                         f"register says '{expected[code]}'")
                    )
                if code in seen:
                    problems.append((label, f"tab '{tab}' repeats code {code}"))
                seen.add(code)

            missing = set(expected) - seen
            if missing:
                sample = ", ".join(sorted(missing)[:6])
                more = "" if len(missing) <= 6 else f" (+{len(missing) - 6} more)"
                problems.append(
                    (label, f"tab '{tab}' is missing {len(missing)}: {sample}{more}")
                )

        if len(per_tab_codes) > 1:
            sets = {t: [c for c, _ in v] for t, v in per_tab_codes.items()}
            first = next(iter(sets))
            for tab, codes in sets.items():
                if codes != sets[first]:
                    problems.append(
                        (label, f"tab '{tab}' rows differ from tab '{first}'")
                    )
                    break

        wb.close()

# ---------------------------------------------------------------- the verdict
print(f"checked {checked} workbooks\n")

if not problems:
    print("PASS - every workbook matches the register.")
    print("       Row counts, codes and names all agree, on both period tabs.")
    sys.exit(0)

by_file = defaultdict(list)
for f, m in problems:
    by_file[f].append(m)

print(f"FAIL - {len(problems)} problem(s) in {len(by_file)} workbook(s):\n")
for f in sorted(by_file):
    print(f"  {f}")
    for m in by_file[f][:4]:
        print(f"      - {m}")
    if len(by_file[f]) > 4:
        print(f"      - ... and {len(by_file[f]) - 4} more")

print(
    "\nIf these are row-count or unknown-code failures, the fix is to regenerate:\n"
    "    cd design\\templates\n"
    "    python generate_templates.py\n"
    "then run this script again."
)
sys.exit(1)
