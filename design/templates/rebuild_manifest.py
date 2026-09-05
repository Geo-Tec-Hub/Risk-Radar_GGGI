#!/usr/bin/env python3
"""
rebuild_manifest.py - rebuild generated/MANIFEST.csv from the workbooks on disk.

WHY THIS EXISTS.  generate_templates.py rewrites MANIFEST.csv on every run
unless --resume is passed.  When the 243 workbooks are regenerated in chunks
(--start/--limit), each chunk therefore overwrites the manifest with only its
own rows, and the file ends up describing the last chunk instead of the set.

The workbooks themselves are complete and correct in that situation - only the
index is short.  This reads the truth back out of the files rather than
re-running the generator, so it costs seconds instead of minutes and cannot
disagree with what is actually on disk.

Every column is taken from the workbook's own _META sheet and its period tab,
except n_with_legacy_weight / n_new_variables, which are counted from the
WEIGHTS sheet the same way the generator counts them: a variable "has a legacy
weight" if the WEIGHTS sheet carries a number for it.

Run it from design/templates:

    python rebuild_manifest.py
"""

import csv
import os
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl is not installed.  Run:  pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
GENERATED = os.path.join(HERE, "generated")
if not os.path.isdir(GENERATED):
    sys.exit(f"No 'generated' folder next to this script ({GENERATED})")

COLUMNS = ["file", "province", "main_sector", "subsector", "hazard",
           "n_variables", "n_with_legacy_weight", "n_new_variables", "n_rows",
           "panel_overrides"]

rows = []
for folder in sorted(os.listdir(GENERATED)):
    fdir = os.path.join(GENERATED, folder)
    if not os.path.isdir(fdir):
        continue
    for fname in sorted(os.listdir(fdir)):
        if not fname.endswith(".xlsx") or fname.startswith("~$"):
            continue
        path = os.path.join(fdir, fname)
        wb = load_workbook(path, read_only=True, data_only=True)

        meta = {}
        if "_META" in wb.sheetnames:
            for r in wb["_META"].iter_rows(values_only=True):
                if r and r[0]:
                    meta[str(r[0]).strip()] = r[1]

        # Data rows start at 5 (row 4 is the locked EXAMPLE row).
        n_rows = 0
        period_tabs = [s for s in wb.sheetnames if "-" in s and s[0].isdigit()]
        if period_tabs:
            ws = wb[period_tabs[0]]
            for r in ws.iter_rows(min_row=5, min_col=1, max_col=1, values_only=True):
                if r[0] is not None and str(r[0]).strip():
                    n_rows += 1

        n_vars = int(meta.get("n_variables") or 0)

        # A variable counts as carrying a legacy weight if WEIGHTS holds a number.
        n_legacy = 0
        if "WEIGHTS" in wb.sheetnames:
            for r in wb["WEIGHTS"].iter_rows(values_only=True):
                for cell in (r[1:] if r else []):
                    if isinstance(cell, (int, float)):
                        n_legacy += 1
                        break

        wb.close()

        rows.append({
            "file": os.path.join("generated", folder, fname).replace("\\", "/"),
            "province": meta.get("province", folder.replace("_", " ")),
            "main_sector": meta.get("main_sector", ""),
            "subsector": meta.get("subsector") or "",
            "hazard": meta.get("hazard", ""),
            "n_variables": n_vars,
            "n_with_legacy_weight": min(n_legacy, n_vars),
            "n_new_variables": max(n_vars - min(n_legacy, n_vars), 0),
            "n_rows": n_rows,
            # written by the generator into _META, e.g. "4 add, 2 retire"
            "panel_overrides": meta.get("panel_overrides", ""),
        })

out = os.path.join(GENERATED, "MANIFEST.csv")
with open(out, "w", newline="", encoding="utf-8") as fh:
    w = csv.DictWriter(fh, fieldnames=COLUMNS)
    w.writeheader()
    w.writerows(rows)

counts = {}
for r in rows:
    counts[r["province"]] = counts.get(r["province"], 0) + 1

print(f"wrote {out}")
print(f"  {len(rows)} workbooks")
for p in sorted(counts):
    n_rows = {r["n_rows"] for r in rows if r["province"] == p}
    print(f"    {p:16s} {counts[p]:3d} workbooks, {sorted(n_rows)} DS rows each")
