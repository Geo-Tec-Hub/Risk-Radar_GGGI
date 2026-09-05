#!/usr/bin/env python3
"""
build_province_overrides.py - derive the province-panel override inputs from a
returned expert review, so the derivation is reproducible rather than hand-typed.

INPUT   the marked-up workbooks a provincial panel returns, e.g.
        "Data sets/CP/CP/Central_want to edit year ranges/" plus its NOTED.xlsx.

OUTPUT  two files read by design/templates/generate_templates.py:
          catalog_addendum.csv           - variable definitions the panel added
                                           that are not yet in FINAL_VARIABLES.xlsx
          province_variable_overrides.csv- per-province add / retire / weight
                                           decisions layered on the national
                                           profile membership

WHY IT WORKS THIS WAY.  FINAL_VARIABLES.xlsx is the experts' own artifact and is
not edited here; the addendum sits beside it and is merged at load time, so what
the panel added is always separable from what the national catalogue says.  The
same applies to membership: the national profile is untouched and the override
file records exactly what Central diverges on and why.

DERIVED WEIGHTS.  Where a combined-period variable is retired in favour of a
25/75 split, the parent's weight is split in the stated proportion so the domain
still totals 100.  Those rows are marked basis="derived" - they are a starting
point for the confirmation screen, NOT an expert decision.

Usage:
    python build_province_overrides.py            # Central, from the paths below
"""
import csv
import datetime
import glob
import os

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
REVIEW = os.path.join(ROOT, "Data sets", "CP", "CP",
                      "Central_want to edit year ranges")
# The build the panel actually received, pinned - NOT generated/, which moves
# every time the generator runs and would silently reduce the diff to nothing.
BASELINE = os.path.join(ROOT, "design", "templates", "baseline_2026-08-15",
                        "Central")
MANIFEST = os.path.join(ROOT, "design", "templates", "generated", "MANIFEST.csv")
PROVINCE = "Central"
PANEL = "Central Province expert panel"
RECEIVED = "2026-09-03"

# Panel's stated split of a combined-period variable into two sub-periods.
# Source: NOTED.xlsx column B, e.g. "No. of Drought events - 1974 to 2004 (25%)".
SPLITS = {
    "DROUGHT_EVENTS_1974_TO_2022": [
        ("DROUGHT_EVENTS_1974_TO_2004", 25), ("DROUGHT_EVENTS_2005_TO_2022", 75)],
    "FLOOD_EVENTS_1974_TO_2023": [
        ("FLOOD_EVENTS_1974_TO_2004", 25), ("FLOOD_EVENTS_2005_TO_2023", 75)],
    "MAXIMUM_DROUGHT_AFFECTED_PEOPLE_DURING_1974_TO_2022": [
        ("MAXIMUM_DROUGHT_AFFECTED_PEOPLE_DURING_1974_TO_2004", 25),
        ("MAXIMUM_DROUGHT_AFFECTED_PEOPLE_DURING_2005_TO_2022", 75)],
    "MAXIMUM_FLOOD_AFFECTED_PEOPLE_DURING_1974_TO_2023": [
        ("MAXIMUM_FLOOD_AFFECTED_PEOPLE_DURING_1974_TO_2004", 25),
        ("MAXIMUM_FLOOD_AFFECTED_PEOPLE_DURING_2005_TO_2023", 75)],
    # NOT a rename. The legacy variable is a COMPOSITE of two things and its own
    # name states the split: "Very wet days (95th percentile) and 3 day
    # cumulative rain fall (out of 35 - 40% for very wet days and 60% for 3 day
    # cumulative respectively)". The panel separated it into its two components,
    # which is why both appear in exactly the same 18 profiles VERY_WET_DAYS did.
    # Read as a 1:1 rename on 2026-09-03 and corrected the same day: that gave
    # very-wet-days the whole weight and left 3-day cumulative unweighted, i.e.
    # excluded from computation - dropping 60% of what the composite measured
    # while the domain still totalled 100, so nothing would have looked wrong.
    "VERY_WET_DAYS": [("VERY_WET_DAYS_95TH_PERCENTILE", 40),
                      ("THREE_DAY_CUMULATIVE_RAINFALL", 60)],
}
# Added with no parent to inherit from - stays unweighted until the panel says.
UNPARENTED = set()

# Codes renamed on the way in. The panel typed the column header by hand and a
# code may not start with a digit: it is not a legal SQL identifier, and the
# importer, the catalogue and every generated column would have to quote it
# forever. The header the panel actually used is kept in the addendum note so
# their returned workbooks can still be matched to the renamed code.
RENAME = {"3DAY_CUMULATIVE_RAINFALL": "THREE_DAY_CUMULATIVE_RAINFALL"}


def profiles():
    """filename -> (main_sector, subsector, hazard) for this province."""
    out = {}
    with open(MANIFEST, encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            if r["province"] == PROVINCE:
                out[os.path.basename(r["file"])] = (
                    r["main_sector"], r["subsector"], r["hazard"])
    return out


def read(path, tab_prefix):
    """{code: (name, domain, column_index)} from a workbook's first data tab."""
    wb = load_workbook(path)
    tab = next(s for s in wb.sheetnames if s.startswith(tab_prefix))
    ws = wb[tab]
    cols = {}
    for c in range(6, ws.max_column + 1):
        code = ws.cell(2, c).value
        if code in (None, "DATA_SOURCE", "NOTES"):
            continue
        sub = (ws.cell(3, c).value or "")
        first = sub.split("\n")[0]
        if ": " in first:
            domain, _, label = first.partition(": ")
        else:
            # a column the panel inserted - it carries a plain label and no
            # domain prefix.  Inherit the domain of the nearest column to its
            # left, which is where the panel placed it.
            domain, label = "", first
        cols[RENAME.get(code, code)] = (label, domain, c)
    last = ""
    for code in list(cols):
        label, domain, c = cols[code]
        if domain:
            last = domain
        else:
            cols[code] = (label, last, c)
    weights = {}
    ws = wb["WEIGHTS"]
    for r in range(4, ws.max_row + 1):
        code = ws.cell(r, 1).value
        if code and not str(code).startswith("SUM"):
            weights[code] = dict(name=ws.cell(r, 2).value, domain=ws.cell(r, 3).value,
                                 rel=ws.cell(r, 4).value, legacy=ws.cell(r, 5).value,
                                 proposed=ws.cell(r, 6).value)
    return cols, weights


def main():
    prof = profiles()
    catalog_rows = {}
    override_rows = []
    seen_defs = {}

    for path in sorted(glob.glob(os.path.join(REVIEW, "*.xlsx"))):
        f = os.path.basename(path)
        if f not in prof:
            continue
        main_sector, subsector, hazard = prof[f]
        new_cols, new_w = read(path, "202")
        old_cols, old_w = read(os.path.join(BASELINE, f), "202")

        added = [c for c in new_cols if c not in old_cols]
        for code in added:
            name, domain, _ = new_cols[code]
            seen_defs.setdefault(code, (name, domain))
            was = next((k for k, v in RENAME.items() if v == code), "")
            catalog_rows[code] = dict(
                variable_code=code, variable_name=name, domain=domain,
                added_by=PANEL, added_on=RECEIVED,
                note="added to the data tab during the Central review; "
                     "not yet in FINAL_VARIABLES.xlsx"
                     + ("; the panel's own column header was '%s' - keep that as "
                        "an indicator_alias so their returned files still match"
                        % was if was else ""))

        # retire a parent only when its stated successor is actually present
        for parent, succ in SPLITS.items():
            if parent not in old_cols:
                continue
            present = [s for s, _ in succ if s in new_cols]
            if not present:
                continue
            override_rows.append(dict(
                province=PROVINCE, main_sector=main_sector, subsector=subsector,
                hazard=hazard, variable_code=parent, action="retire",
                weight_pct="", basis="",
                note="superseded by " + " + ".join(present)))
            base = old_w.get(parent, {}).get("legacy")
            for code, share in succ:
                if code not in new_cols:
                    continue
                w = ""
                basis = ""
                if base not in (None, ""):
                    w = round(float(base) * share / 100.0, 3)
                    w = int(w) if float(w).is_integer() else w
                    basis = ("derived: %d%% of retired %s (%s)"
                             % (share, parent, base))
                override_rows.append(dict(
                    province=PROVINCE, main_sector=main_sector,
                    subsector=subsector, hazard=hazard, variable_code=code,
                    action="add", weight_pct=w, basis=basis,
                    note="replaces " + parent))

        for code in added:
            if code in UNPARENTED:
                override_rows.append(dict(
                    province=PROVINCE, main_sector=main_sector,
                    subsector=subsector, hazard=hazard, variable_code=code,
                    action="add", weight_pct="", basis="",
                    note="new variable; unweighted until the panel confirms"))

        # weights the panel typed into the PROPOSED column
        for code, meta in new_w.items():
            if code in old_w and old_w[code]["proposed"] != meta["proposed"] \
                    and meta["proposed"] not in (None, ""):
                override_rows.append(dict(
                    province=PROVINCE, main_sector=main_sector,
                    subsector=subsector, hazard=hazard, variable_code=code,
                    action="weight", weight_pct=meta["proposed"],
                    basis="panel: entered in the returned WEIGHTS tab",
                    note="was unweighted"))

    with open(os.path.join(HERE, "catalog_addendum.csv"), "w", newline="",
              encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["variable_code", "variable_name",
                                           "domain", "added_by", "added_on", "note"])
        w.writeheader()
        for k in sorted(catalog_rows):
            w.writerow(catalog_rows[k])

    with open(os.path.join(HERE, "province_variable_overrides.csv"), "w",
              newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["province", "main_sector", "subsector",
                                           "hazard", "variable_code", "action",
                                           "weight_pct", "basis", "note"])
        w.writeheader()
        for r in sorted(override_rows, key=lambda r: (r["main_sector"], r["subsector"],
                                                      r["hazard"], r["action"],
                                                      r["variable_code"])):
            w.writerow(r)

    print("catalog_addendum.csv            %3d new variable definitions"
          % len(catalog_rows))
    for a in ("add", "retire", "weight"):
        print("province_variable_overrides.csv %3d %s rows"
              % (sum(1 for r in override_rows if r["action"] == a), a))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
