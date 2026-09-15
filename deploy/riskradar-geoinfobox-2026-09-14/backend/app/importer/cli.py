"""Dry-run entrypoint: parse a province workbook, resolve every column and
DS_Code, print what it found and what it would reject. Writes nothing.

    python -m app.importer.cli "<path to workbook.xlsx>" [--sheet NAME] [--province Western]

Tries the live database first for the indicator catalogue/aliases; falls back
to the committed generated snapshot if it's unreachable (see catalog.py). The
division register always comes from dsd_register.csv -- that's the only place
carrying the numeric DS_Code the workbooks use.
"""

from __future__ import annotations

import argparse
import asyncio
import sys
from pathlib import Path

import openpyxl

from app.importer.catalog import Catalog, load_catalog
from app.importer.reader import parse_sheet
from app.importer.report import render_sheet_report
from app.importer.validate import validate_sheet

# Sheet-name prefixes observed across the province workbooks (brief: "Province
# comes from the sheet name or the filename").
_PREFIX_TO_PROVINCE = {
    "wp": "Western",
    "sp": "Southern",
    "cp": "Central",
    "nw": "Northwestern",
    "sg": "Sabaragamuwa",
    "ep": "Eastern",
    "nc": "North Central",
    "np": "Northern",
    "uva": "Uva",
}

_KNOWN_PROVINCES = [
    "Central", "Eastern", "North Central", "Northern",
    "Northwestern", "Sabaragamuwa", "Southern", "Uva", "Western",
]


def guess_province(sheet_name: str, path: Path) -> str | None:
    prefix = sheet_name.split("_")[0].lower()
    if prefix in _PREFIX_TO_PROVINCE:
        return _PREFIX_TO_PROVINCE[prefix]
    haystack = " / ".join(path.parts).lower()
    for prov in _KNOWN_PROVINCES:
        if prov.lower() in haystack:
            return prov
    return None


async def _load_catalog_best_effort() -> Catalog:
    pool = None
    try:
        from app.db import create_pool

        pool = await asyncio.wait_for(create_pool(), timeout=5)
    except Exception:
        pool = None
    catalog = await load_catalog(pool)
    if pool is not None:
        await pool.close()
    return catalog


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("workbook", type=Path, help="path to the .xlsx workbook")
    parser.add_argument("--sheet", help="only parse this sheet (default: every sheet)")
    parser.add_argument("--province", help="expected province (default: guessed from sheet name / path)")
    args = parser.parse_args(argv)

    if not args.workbook.exists():
        print(f"file not found: {args.workbook}", file=sys.stderr)
        return 2

    catalog = asyncio.run(_load_catalog_best_effort())
    print(f"catalogue/register source: {catalog.source}")
    print(f"registered divisions: {len(catalog.divisions_by_num)}")
    print()

    wb = openpyxl.load_workbook(args.workbook, data_only=True, read_only=True)
    sheet_names = [args.sheet] if args.sheet else wb.sheetnames

    any_rejected = False
    for name in sheet_names:
        if name not in wb.sheetnames:
            print(f"=== Sheet '{name}' === NOT FOUND in this workbook (has: {', '.join(wb.sheetnames)})")
            any_rejected = True
            continue
        ws = wb[name]
        parsed = parse_sheet(ws, catalog)
        expected_province = args.province or guess_province(name, args.workbook)
        report = validate_sheet(parsed, catalog, expected_province)
        print(render_sheet_report(report))
        print()
        any_rejected = any_rejected or not report.would_load

    return 1 if any_rejected else 0


if __name__ == "__main__":
    raise SystemExit(main())
