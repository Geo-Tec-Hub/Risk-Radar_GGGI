"""Strict reader for the GENERATED upload templates (design/templates).

WHY THIS IS A SECOND READER.  `reader.py` parses the legacy nine-province
workbooks: multi-block sheets, zone label rows, a numeric DS_Code and free-text
variable labels that have to be reconciled against the catalogue. Those files
are human artefacts and the reader's job is to survive them.

These are different files. The system generated them, so their shape is known
exactly and a surprise is a defect rather than something to accommodate. That is
the dual-path ingestion settled in D3 rev 5: a STRICT path for templates and a
reconciliation wizard for legacy sheets. Trying to serve both from one reader
would mean the strict path inherits the lenient path's guesses.

THE CONTRACT IS `_META`, NOT THE COLUMN HEADINGS.  The hidden `_META` sheet
carries `profile_code`, `expected_columns` and `periods`. Validating against it
means a renamed, reordered or inserted column is caught as a structural error
naming the file, instead of being silently skipped.

CLIMATE WORKBOOKS ARE A SECOND KIND, DECLARED IN `_META`. A workbook with
`kind = climate` carries the hazard-domain variables for a whole province and
belongs to no sector, so its `_META` has no main_sector/subsector/hazard to
cross-check. `kind` defaults to "profile" so every workbook generated before
this existed keeps its current meaning.

THE `WEIGHTS` TAB IS READ, AND IS ADVISORY. The generated templates carry a
WEIGHTS tab whose instruction line has always promised that "the importer READS
this tab and pre-fills the confirmation screen". Until now nothing read it, so
a panel could fill the yellow PROPOSED column and have it silently go nowhere.
It is read here and returned alongside the values, but it is NEVER written:
`save_profile_weights()` stays the one audited path that changes a weight, and
what this reader produces is a proposal for a person to confirm there. A missing
or malformed WEIGHTS tab is a warning, never an error -- the values in the
period tabs are the import, and a bad reference tab must not refuse them.

PERIOD COMES FROM THE TAB NAME.  Not from the YEAR_START/YEAR_END cells. In the
returned Central files those cells still read 2020/2025 because they were locked
when the panel renamed the tabs to 2021-2025 / 2026-2030 (their NOTED.xlsx says
so). Milinda's call, 3 Sep 2026: the tab is authoritative, the cells are
disregarded. The reader records the disagreement as a warning so it is visible,
and refuses only if the tab name itself is unreadable.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Optional

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

FIXED_COLUMNS = ["DS_CODE", "DS_DIVISION", "DISTRICT", "YEAR_START", "YEAR_END"]
TRAILING_COLUMNS = ["DATA_SOURCE", "NOTES"]
HEADER_ROW = 2
EXAMPLE_ROW = 4          # the grey specimen row; never a data row
FIRST_DATA_ROW = 5
TAB_RE = re.compile(r"^(\d{4})-(\d{4})$")

# The WEIGHTS tab, as generate_templates.fill_weights_sheet lays it out:
# title, instruction block, header at row 3, data from row 4 until the first
# blank code, then a spacer and the "SUM - <domain>" rows.
WEIGHTS_SHEET = "WEIGHTS"
WEIGHTS_HEADER_ROW = 3
WEIGHTS_FIRST_ROW = 4
WEIGHTS_SUM_RE = re.compile(r"^\s*SUM\s*-", re.I)


@dataclass(frozen=True)
class TemplateValue:
    excel_row: int
    ds_code: str
    variable_code: str
    raw_value: float
    data_source: Optional[str]
    notes: Optional[str]


@dataclass(frozen=True)
class TemplateWeight:
    """One row of the WEIGHTS tab.

    `proposed_pct` is the panel's yellow column -- what they are asking for.
    `legacy_pct` is what the previous provincial workbook carried, printed by
    the generator for reference; blank means the variable is new. Both are
    optional because a blank weight is a real statement (not yet decided) and
    must never be read as a zero."""
    excel_row: int
    variable_code: str
    variable_name: Optional[str] = None
    domain: Optional[str] = None
    relationship: Optional[str] = None
    legacy_pct: Optional[float] = None
    proposed_pct: Optional[float] = None
    basis: Optional[str] = None


@dataclass
class TemplateSheet:
    tab: str
    year_start: int
    year_end: int
    values: list[TemplateValue] = field(default_factory=list)
    rows_scanned: int = 0
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class TemplateWorkbook:
    filename: str
    profile_code: Optional[str] = None
    kind: str = "profile"          # "profile" | "climate"
    province: Optional[str] = None
    main_sector: Optional[str] = None
    subsector: Optional[str] = None
    hazard: Optional[str] = None
    protection: Optional[str] = None
    expected_columns: list[str] = field(default_factory=list)
    sheets: list[TemplateSheet] = field(default_factory=list)
    weights: list[TemplateWeight] = field(default_factory=list)
    weights_tab_present: bool = False
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def would_load(self) -> bool:
        return not self.errors and not any(s.errors for s in self.sheets)

    @property
    def value_count(self) -> int:
        return sum(len(s.values) for s in self.sheets)


def read_meta(wb: openpyxl.Workbook) -> dict[str, str]:
    if "_META" not in wb.sheetnames:
        return {}
    ws = wb["_META"]
    out: dict[str, str] = {}
    for r in range(1, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if k:
            out[str(k).strip()] = ws.cell(r, 2).value
    return out


def _header(ws: Worksheet) -> list[str]:
    return [(ws.cell(HEADER_ROW, c).value or "") and str(ws.cell(HEADER_ROW, c).value).strip()
            for c in range(1, ws.max_column + 1)]


def read_sheet(ws: Worksheet, expected: list[str],
               aliases: Optional[dict[str, str]] = None,
               ordered: bool = True,
               signed: Optional[set[str]] = None) -> TemplateSheet:
    m = TAB_RE.match(ws.title)
    if not m:
        s = TemplateSheet(tab=ws.title, year_start=0, year_end=0)
        s.errors.append("tab name %r is not a YYYY-YYYY period" % ws.title)
        return s
    y0, y1 = int(m.group(1)), int(m.group(2))
    sheet = TemplateSheet(tab=ws.title, year_start=y0, year_end=y1)

    header = [(aliases or {}).get(h, h) for h in _header(ws)]
    if expected and header != expected:
        missing = [c for c in expected if c not in header]
        extra = [c for c in header if c and c not in expected]
        dupes = sorted({c for c in header if c and header.count(c) > 1})
        # A column the contract has and the file does not carries no data, so in
        # review-copy mode it is reported and skipped rather than failing the
        # file -- the panel was invited to drop columns. An UNKNOWN column is
        # always an error: it means a value nothing can be keyed to.
        if missing:
            (sheet.errors if ordered else sheet.warnings).append(
                "columns absent from the file: " + ", ".join(missing))
        if extra:
            sheet.errors.append("columns not in the contract: " + ", ".join(extra))
        if dupes:
            sheet.errors.append(
                "two columns resolve to the same variable: " + ", ".join(dupes))
        if ordered and not missing and not extra:
            sheet.errors.append("columns are in a different order than _META states")
    if sheet.errors:
        return sheet

    idx = {name: i + 1 for i, name in enumerate(header) if name}
    var_cols = [(name, col) for name, col in idx.items()
                if name not in FIXED_COLUMNS and name not in TRAILING_COLUMNS]

    mismatched_years = 0
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        ds_code = ws.cell(r, idx["DS_CODE"]).value
        if ds_code in (None, ""):
            continue
        sheet.rows_scanned += 1
        cell_y0 = ws.cell(r, idx["YEAR_START"]).value
        cell_y1 = ws.cell(r, idx["YEAR_END"]).value
        if (cell_y0, cell_y1) != (y0, y1):
            mismatched_years += 1
        src = ws.cell(r, idx["DATA_SOURCE"]).value if "DATA_SOURCE" in idx else None
        note = ws.cell(r, idx["NOTES"]).value if "NOTES" in idx else None
        for name, col in var_cols:
            v = ws.cell(r, col).value
            if v in (None, ""):
                continue          # blank means not collected - never a zero
            if isinstance(v, str):
                sheet.errors.append(
                    "row %d, column %s: %r is text, not a number" % (r, name, v[:40]))
                continue
            if float(v) < 0 and name not in (signed or set()):
                # Per-variable, never global: a variable declared 'signed' in
                # indicator_catalog expresses a change or trend, where negative
                # is meaningful. Everything else is a quantity and a negative
                # one is a data error worth stopping the file for.
                sheet.errors.append(
                    "row %d, column %s: %s is negative, and %s is not declared "
                    "as a signed variable" % (r, name, v, name))
                continue
            sheet.values.append(TemplateValue(
                excel_row=r, ds_code=str(ds_code).strip(), variable_code=name,
                raw_value=float(v),
                data_source=str(src).strip() if src else None,
                notes=str(note).strip() if note else None))

    if mismatched_years:
        # Not an error: the tab is authoritative (3 Sep 2026). Reported so a
        # stale year cell is never silently absorbed.
        sheet.warnings.append(
            "%d of %d rows carry YEAR_START/YEAR_END %s-%s, disagreeing with the "
            "tab; the tab name wins and %d-%d was used"
            % (mismatched_years, sheet.rows_scanned, cell_y0, cell_y1, y0, y1))
    return sheet


def _num(v) -> tuple[Optional[float], bool]:
    """(value, ok). Blank is (None, True) -- an undecided weight, not a zero.
    A formula string is (None, True) too: the file was opened with data_only,
    so a string here means the cached result was never written by Excel."""
    if v in (None, ""):
        return None, True
    if isinstance(v, str):
        t = v.strip().replace("%", "")
        if not t or t.startswith("="):
            return None, True
        try:
            return float(t), True
        except ValueError:
            return None, False
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None, False
    # The generator writes 0-100, but a hand-edited cell formatted as a
    # percentage comes back as 0.25 for 25%. Left alone rather than guessed at:
    # 0.25 is a legal weight, and silently multiplying it by 100 would be the
    # importer inventing a number. The range check below flags what looks off.
    return f, True


def read_weights(ws: Worksheet,
                 aliases: Optional[dict[str, str]] = None
                 ) -> tuple[list[TemplateWeight], list[str]]:
    """Read the WEIGHTS tab. Returns (rows, warnings) -- never errors.

    Stops at the first blank variable_code, which is the spacer the generator
    puts before the SUM rows; the SUM rows are skipped explicitly as well so a
    hand-edited file that closed that gap still reads correctly.
    """
    rows: list[TemplateWeight] = []
    warnings: list[str] = []
    alias_map = aliases or {}

    header = [(ws.cell(WEIGHTS_HEADER_ROW, c).value or "") and
              str(ws.cell(WEIGHTS_HEADER_ROW, c).value).strip().lower()
              for c in range(1, min(ws.max_column, 8) + 1)]
    if not header or "variable_code" not in (header[0] or ""):
        warnings.append(
            "the WEIGHTS tab does not have `variable_code` in the first column "
            "of row %d, so it was not read. The values in the period tabs were "
            "imported as normal." % WEIGHTS_HEADER_ROW)
        return rows, warnings

    seen: set[str] = set()
    for r in range(WEIGHTS_FIRST_ROW, ws.max_row + 1):
        raw_code = ws.cell(r, 1).value
        if raw_code in (None, ""):
            break
        code = str(raw_code).strip()
        if WEIGHTS_SUM_RE.match(code):
            continue
        code = alias_map.get(code, code)
        if code in seen:
            warnings.append("WEIGHTS row %d: %s appears twice; the first row was used"
                            % (r, code))
            continue
        seen.add(code)

        name = ws.cell(r, 2).value
        domain = ws.cell(r, 3).value
        rel = ws.cell(r, 4).value
        legacy, legacy_ok = _num(ws.cell(r, 5).value)
        proposed, proposed_ok = _num(ws.cell(r, 6).value)
        basis = ws.cell(r, 7).value

        if not proposed_ok:
            warnings.append("WEIGHTS row %d (%s): the proposed weight %r is not a "
                            "number and was ignored"
                            % (r, code, str(ws.cell(r, 6).value)[:30]))
        elif proposed is not None and not (0 <= proposed <= 100):
            warnings.append("WEIGHTS row %d (%s): the proposed weight %s is outside "
                            "0-100" % (r, code, proposed))
        if not legacy_ok:
            warnings.append("WEIGHTS row %d (%s): the legacy weight is not a number "
                            "and was ignored" % (r, code))

        rows.append(TemplateWeight(
            excel_row=r, variable_code=code,
            variable_name=str(name).strip() if name else None,
            domain=str(domain).strip().lower() if domain else None,
            relationship=str(rel).strip() if rel else None,
            legacy_pct=legacy, proposed_pct=proposed,
            basis=str(basis).strip() if basis else None))

    # Totals are reported, never enforced here. An incomplete weights tab is a
    # normal mid-panel state and must not colour the import's verdict; the
    # confirmation screen is where 100 has to be reached.
    by_domain: dict[str, float] = {}
    filled: dict[str, int] = {}
    for w in rows:
        if w.proposed_pct is None or not w.domain:
            continue
        by_domain[w.domain] = by_domain.get(w.domain, 0.0) + w.proposed_pct
        filled[w.domain] = filled.get(w.domain, 0) + 1
    for dom, total in sorted(by_domain.items()):
        if abs(total - 100.0) > 0.01:
            warnings.append(
                "WEIGHTS: the proposed %s weights total %g over %d variable(s), "
                "not 100. Nothing is blocked by this -- weights take effect only "
                "when they are confirmed in the weights editor."
                % (dom, round(total, 4), filled[dom]))
    if rows and not by_domain:
        warnings.append(
            "WEIGHTS: the tab was read but no proposed weight was filled in, so "
            "there is nothing to confirm.")
    return rows, warnings


def read_workbook(path: str, expected: Optional[list[str]] = None,
                  aliases: Optional[dict[str, str]] = None,
                  signed: Optional[set[str]] = None) -> TemplateWorkbook:
    """Read one generated upload template.

    `expected` overrides the `_META` column contract. Pass it ONLY for a review
    copy the panel was invited to restructure: their `_META` still describes the
    build they were sent, so validating against it would reject the very changes
    that were asked for. For a locked issue leave it None, so `_META` governs
    and a changed structure is an error.

    `aliases` maps a column heading to the catalogue code it means - the
    `indicator_alias` case, e.g. a panel's hand-typed `3DAY_CUMULATIVE_RAINFALL`
    for `THREE_DAY_CUMULATIVE_RAINFALL`.
    """
    out = TemplateWorkbook(filename=path.replace("\\", "/").split("/")[-1])
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        # A .xlsx is a zip. Anything that is not one - a CSV renamed to .xlsx,
        # a truncated download, an .xls saved under the wrong extension -
        # raises out of openpyxl, and before 4 Sep 2026 that propagated all the
        # way to a 500. An upload the system cannot read is a REFUSAL with a
        # reason, not a server error: a 500 tells the person nothing about what
        # they did wrong and reads as "the site is broken".
        out.errors.append(
            "this file could not be opened as an .xlsx workbook (%s: %s). "
            "A file renamed to .xlsx is still not one - export it from Excel "
            "as .xlsx, or re-download it."
            % (type(exc).__name__, str(exc)[:120]))
        return out
    meta = read_meta(wb)
    if not meta:
        out.errors.append("no _META sheet - this is not a generated upload template")
        return out

    out.profile_code = meta.get("profile_code")
    # Absent on every workbook generated before climate templates existed, and
    # those are all sector files -- hence the default rather than an error.
    out.kind = str(meta.get("kind") or "profile").strip().lower()
    out.province = meta.get("province")
    out.main_sector = meta.get("main_sector")
    out.subsector = meta.get("subsector") or None
    out.hazard = meta.get("hazard")
    out.protection = meta.get("protection")
    out.expected_columns = [c for c in str(meta.get("expected_columns") or "").split("|") if c]
    contract = expected if expected is not None else out.expected_columns
    if expected is not None:
        out.warnings.append(
            "columns validated against the CURRENT profile, not this file's _META "
            "- review-copy mode")

    if out.protection == "unlocked-review":
        out.warnings.append(
            "this is an unlocked REVIEW copy - its structure was open for editing, "
            "so the _META contract is a weaker guarantee than on a locked issue")

    for name in wb.sheetnames:
        if TAB_RE.match(name):
            out.sheets.append(read_sheet(wb[name], contract, aliases,
                                         ordered=expected is None,
                                         signed=signed))
    if not out.sheets:
        out.errors.append("no period tabs found")

    # Advisory, and last: whatever it says, it cannot change whether the file
    # loads. A workbook generated before the WEIGHTS tab existed simply has none.
    if WEIGHTS_SHEET in wb.sheetnames:
        out.weights_tab_present = True
        out.weights, w_warn = read_weights(wb[WEIGHTS_SHEET], aliases)
        out.warnings.extend(w_warn)
    else:
        out.warnings.append(
            "this workbook has no WEIGHTS tab, so no weights were proposed with "
            "it. Weights are set in the weights editor.")
    return out
