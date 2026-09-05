"""Offline tests for the Stage 3 importer's reader + validator
(TASK_BRIEF_IMPORTER.md). No database needed: synthetic workbooks exercise
the rules one at a time, and two tests run the real reader against real
workbooks under `Data sets/` using the static catalogue snapshot.

Run: backend/venv/Scripts/python -m pytest backend/tests/test_importer_reader.py -v
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app.importer.catalog import Catalog, CatalogEntry, Division, load_static_snapshot, normalize_header
from app.importer.reader import ColumnKind, parse_sheet
from app.importer.validate import validate_sheet

REPO_ROOT = Path(__file__).resolve().parents[2]


# ---------------------------------------------------------------------------
# A small synthetic catalogue: 2 divisions in "Testland", 2 indicators.
# ---------------------------------------------------------------------------
def make_catalog() -> Catalog:
    divisions = [
        Division(ds_code_num=1, code="TL1", name="Alpha", province="Testland", district="Test District"),
        Division(ds_code_num=2, code="TL2", name="Beta", province="Testland", district="Test District"),
    ]
    entries = [
        CatalogEntry(code="RICE_EXTENT", name="Rice extent", status="active", domain="exposure"),
        CatalogEntry(code="RICE_FARMERS", name="Rice farmers", status="active", domain="exposure"),
        CatalogEntry(code="RICE_EXTENT_FINAL", name="Rice extent (combined)", status="active", domain="hazard"),
        CatalogEntry(code="STRONG_WIND_EVENTS", name="Strong wind events", status="active", domain="hazard"),
    ]
    return Catalog(
        source="synthetic test catalogue",
        divisions_by_num={d.ds_code_num: d for d in divisions},
        divisions_by_province={"testland": divisions},
        indicator_by_norm_code={normalize_header(e.code): e for e in entries},
        alias_by_norm={normalize_header("Paddy area"): entries[0]},
    )


def sheet_with_rows(header_row_index: int, headers: list[str], rows: list[list]):
    """Build a worksheet with `headers` at row `header_row_index` (1-based),
    preceded by blank rows, followed by `rows` of data."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for _ in range(header_row_index - 1):
        ws.append([])
    ws.append(headers)
    for row in rows:
        ws.append(row)
    return ws


IDENTITY = ["PROVINCE_N", "DISTRICT_N", "DSD_N", "DS_Code"]


def test_normalize_header_strips_and_lowercases():
    assert normalize_header("DS_Code") == normalize_header("DS Code") == "dscode"
    assert normalize_header("PCT_PADDY_LAND_TOTAL _LAND_AREA") == normalize_header("PCT_PADDY_LAND_TOTAL_LAND_AREA")


def test_header_row_found_beyond_row_one():
    # Header on row 4, per the brief's "7.Paddy_Subsector puts it on row 4" example.
    ws = sheet_with_rows(4, IDENTITY + ["RICE_EXTENT"], [["Testland", "Test District", "Alpha", 1, 10]])
    parsed = parse_sheet(ws, make_catalog())
    assert parsed.header_row == 4
    assert len(parsed.rows) == 1


def test_header_not_found_rejects_sheet():
    # No cell normalising to 'dscode' anywhere in the first 10 rows.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Variable", "Weight"])
    ws.append(["H1", 70])
    report = validate_sheet(parse_sheet(ws, make_catalog()), make_catalog())
    assert not report.would_load
    assert report.sheet_findings[0].code == "NO_HEADER"


def test_blank_header_columns_are_ignored_not_rejected():
    ws = sheet_with_rows(
        1,
        IDENTITY + ["RICE_EXTENT", None, None],
        [["Testland", "Test District", "Alpha", 1, 10, "No Paddy", None]],
    )
    parsed = parse_sheet(ws, make_catalog())
    # only the 5 named columns are classified -- the two blank-header columns never appear
    assert len(parsed.columns) == 5
    report = validate_sheet(parsed, make_catalog())
    assert report.would_load


def test_literal_normalization_calculation_headers_are_skipped():
    # A column literally headed 'Normalization' / '% Calculation' (no separate
    # zone-label row needed) is still recognised via the own-header fallback.
    ws = sheet_with_rows(
        1,
        IDENTITY + ["RICE_EXTENT", "Normalization", "% Calculation"],
        [["Testland", "Test District", "Alpha", 1, 10, 0.5, 35]],
    )
    parsed = parse_sheet(ws, make_catalog())
    kinds = {c.raw_header: c.kind.value for c in parsed.columns}
    assert kinds["Normalization"] == "ignored_computed"
    assert kinds["% Calculation"] == "ignored_computed"
    report = validate_sheet(parsed, make_catalog())
    assert report.would_load
    assert len(report.accepted_values) == 1
    assert report.accepted_values[0].indicator.code == "RICE_EXTENT"


def _sheet_with_zone_rows(zone_row2, zone_row3, headers, data_rows):
    """Mirrors the real hazard files: block-group labels split across the two
    rows directly above the header row (row 4), metadata on row 1 unused."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([])
    ws.append(zone_row2)
    ws.append(zone_row3)
    ws.append(headers)
    for row in data_rows:
        ws.append(row)
    return ws


def test_zone_labelled_normalization_block_ignored_regardless_of_its_own_header_text():
    # F's own header doesn't look like 'Normalization' at all -- only the zone
    # label (row 3, forward-filled) says so. Mirrors Drought_Hazard_Data's
    # 'SPI (sevirity)' under a 'Normalization' zone.
    headers = IDENTITY + ["RICE_EXTENT", "Rice extent normalised value"]
    zone_row3 = [None] * 4 + [None, "Normalization"]
    ws = _sheet_with_zone_rows([None] * 6, zone_row3, headers, [["Testland", "Test District", "Alpha", 1, 10, 0.4]])
    parsed = parse_sheet(ws, make_catalog())
    col = next(c for c in parsed.columns if c.raw_header == "Rice extent normalised value")
    assert col.kind == ColumnKind.IGNORED_COMPUTED
    assert "Normalization" in (col.note or "")
    report = validate_sheet(parsed, make_catalog())
    assert report.would_load
    assert len(report.accepted_values) == 1  # only RICE_EXTENT, not its normalised copy


def test_web_column_preceded_by_percent_weighted_column_is_a_frozen_composite():
    # Mirrors Sheet6 / the *_EVENTS_* composites: a %Calculation block ending
    # in an explicit '(75%)'-style column, immediately followed by a WEB
    # column whose header IS a real catalogue code -- still excluded, since
    # it's the Excel-frozen combination the brief says never to import.
    headers = IDENTITY + ["RICE_EXTENT", "Rice extent - early (25%)", "Rice extent - late (75%)", "RICE_EXTENT_FINAL"]
    zone_row3 = [None] * 4 + [None, "% Calculation", None, "WEB"]
    ws = _sheet_with_zone_rows([None] * 8, zone_row3, headers, [["Testland", "Test District", "Alpha", 1, 10, 2, 8, 10]])
    parsed = parse_sheet(ws, make_catalog())
    col = next(c for c in parsed.columns if c.raw_header == "RICE_EXTENT_FINAL")
    assert col.kind == ColumnKind.IGNORED_COMPUTED
    assert "frozen composite" in (col.note or "")
    report = validate_sheet(parsed, make_catalog())
    assert all(av.indicator.code != "RICE_EXTENT_FINAL" for av in report.accepted_values)


def test_web_column_not_preceded_by_percent_weighting_is_kept_as_raw():
    # Mirrors STRONG_WIND_EVENTS / LIGHTNING_INCIDENTS in Flood_Hazard_Data:
    # labelled 'WEB' too, but it's the sheet's only column for that variable
    # (no period-pair, no % split before it) -- a raw value, not a composite.
    headers = IDENTITY + ["RICE_EXTENT", "STRONG_WIND_EVENTS"]
    zone_row3 = [None] * 4 + [None, "WEB"]
    ws = _sheet_with_zone_rows([None] * 6, zone_row3, headers, [["Testland", "Test District", "Alpha", 1, 10, 3]])
    parsed = parse_sheet(ws, make_catalog())
    col = next(c for c in parsed.columns if c.raw_header == "STRONG_WIND_EVENTS")
    assert col.kind == ColumnKind.INDICATOR
    assert col.indicator.code == "STRONG_WIND_EVENTS"
    report = validate_sheet(parsed, make_catalog())
    assert report.would_load
    assert any(av.indicator.code == "STRONG_WIND_EVENTS" for av in report.accepted_values)


def test_duplicate_of_an_unresolved_column_is_ignored_not_independently_unresolved():
    # Mirrors Sheet6's H/I: an exact duplicate of an earlier (unresolved)
    # header, with no zone label covering it at all -- still only one
    # finding, not two, for what is structurally one block copy.
    headers = IDENTITY + ["Some prose header", "Some prose header"]
    ws = sheet_with_rows(1, headers, [["Testland", "Test District", "Alpha", 1, 5, 5]])
    parsed = parse_sheet(ws, make_catalog())
    kinds = [c.kind for c in parsed.columns if c.raw_header == "Some prose header"]
    assert kinds == [ColumnKind.UNRESOLVED, ColumnKind.IGNORED_COMPUTED]
    report = validate_sheet(parsed, make_catalog())
    assert len([f for f in report.sheet_findings if f.code == "V011"]) == 1


def test_unresolved_column_rejects_the_sheet():
    ws = sheet_with_rows(
        1,
        IDENTITY + ["Some Prose Header Nobody Mapped"],
        [["Testland", "Test District", "Alpha", 1, 10]],
    )
    report = validate_sheet(parse_sheet(ws, make_catalog()), make_catalog())
    assert not report.would_load
    assert any(f.code == "V011" for f in report.sheet_findings)


def test_alias_resolves_a_prose_header():
    ws = sheet_with_rows(1, IDENTITY + ["Paddy area"], [["Testland", "Test District", "Alpha", 1, 10]])
    parsed = parse_sheet(ws, make_catalog())
    col = next(c for c in parsed.columns if c.raw_header == "Paddy area")
    assert col.indicator.code == "RICE_EXTENT"


def test_duplicate_ds_code_rejects_the_sheet():
    ws = sheet_with_rows(
        1,
        IDENTITY + ["RICE_EXTENT"],
        [
            ["Testland", "Test District", "Alpha", 1, 10],
            ["Testland", "Test District", "Alpha", 1, 11],
        ],
    )
    report = validate_sheet(parse_sheet(ws, make_catalog()), make_catalog())
    assert not report.would_load
    assert any(f.code == "V002" for f in report.sheet_findings)


def test_unknown_ds_code_is_a_row_finding_not_a_crash():
    ws = sheet_with_rows(
        1,
        IDENTITY + ["RICE_EXTENT"],
        [
            ["Testland", "Test District", "Alpha", 1, 10],
            ["Testland", "Test District", "Ghost", 999, 5],
        ],
    )
    report = validate_sheet(parse_sheet(ws, make_catalog()), make_catalog())
    assert not report.would_load
    assert any(f.code == "V003" and "999" in f.message for f in report.row_findings)
    # the good row still gets validated and would have loaded on its own
    assert len(report.accepted_values) == 1


def test_blank_ds_code_row_rejected_as_v001():
    ws = sheet_with_rows(1, IDENTITY + ["RICE_EXTENT"], [["Testland", "Test District", "Alpha", None, 10]])
    report = validate_sheet(parse_sheet(ws, make_catalog()), make_catalog())
    assert any(f.code == "V001" for f in report.row_findings)


def test_non_numeric_cell_rejected_as_v008():
    ws = sheet_with_rows(1, IDENTITY + ["RICE_EXTENT"], [["Testland", "Test District", "Alpha", 1, "lots"]])
    report = validate_sheet(parse_sheet(ws, make_catalog()), make_catalog())
    assert not report.would_load
    assert any(f.code == "V008" and "lots" in f.message for f in report.cell_findings)


def test_excel_error_treated_as_absent_never_as_zero():
    ws = sheet_with_rows(1, IDENTITY + ["RICE_EXTENT"], [["Testland", "Test District", "Alpha", 1, "#DIV/0!"]])
    report = validate_sheet(parse_sheet(ws, make_catalog()), make_catalog())
    assert report.would_load  # an absent value is not an error
    assert report.excel_error_cells == 1
    assert report.accepted_values == []  # never fabricated as 0


def test_blank_cell_is_absent_not_zero():
    ws = sheet_with_rows(1, IDENTITY + ["RICE_EXTENT"], [["Testland", "Test District", "Alpha", 1, None]])
    report = validate_sheet(parse_sheet(ws, make_catalog()), make_catalog())
    assert report.would_load
    assert report.accepted_values == []


def test_province_mismatch_rejects_the_sheet():
    ws = sheet_with_rows(1, IDENTITY + ["RICE_EXTENT"], [["Nottestland", "Test District", "Alpha", 1, 10]])
    report = validate_sheet(parse_sheet(ws, make_catalog()), make_catalog(), expected_province="Testland")
    assert not report.would_load
    assert any(f.code == "PROVINCE_MISMATCH" for f in report.sheet_findings)


def test_outstanding_division_reported_when_missing_from_sheet():
    # Only division 1 (Alpha) appears; division 2 (Beta) is registered but absent.
    ws = sheet_with_rows(1, IDENTITY + ["RICE_EXTENT"], [["Testland", "Test District", "Alpha", 1, 10]])
    report = validate_sheet(parse_sheet(ws, make_catalog()), make_catalog(), expected_province="Testland")
    assert report.would_load
    assert [d.code for d in report.outstanding_divisions] == ["TL2"]


def test_happy_path_loads_cleanly():
    ws = sheet_with_rows(
        1,
        IDENTITY + ["RICE_EXTENT", "RICE_FARMERS"],
        [
            ["Testland", "Test District", "Alpha", 1, 10, 100],
            ["Testland", "Test District", "Beta", 2, 20, 200],
        ],
    )
    report = validate_sheet(parse_sheet(ws, make_catalog()), make_catalog(), expected_province="Testland")
    assert report.would_load
    assert len(report.accepted_values) == 4
    assert report.outstanding_divisions == []


# ---------------------------------------------------------------------------
# Integration: the real catalogue snapshot against real workbooks under
# `Data sets/`. The hazard-block workbook named in the brief
# (Paddy_Drought 1.xlsx / Final_Drought oracle) is not present in this repo
# checkout, so these exercise what IS available: real header-row variation,
# real DS_Code joins against the 340-division register, and the real prose
# headers that need aliases.
# ---------------------------------------------------------------------------
WESTERN_PADDY = REPO_ROOT / "Data sets" / "Data sets" / "7.Paddy_Subsector.xlsx"
MAX_AFFECTED = REPO_ROOT / "Data sets" / "Data sets" / "3.Flood_Drought_Lslide-maximum_affected_data.xlsx"


@pytest.mark.skipif(not WESTERN_PADDY.exists(), reason="Data sets/ workbook not present in this checkout")
def test_real_western_paddy_sheet_loads_cleanly():
    catalog = load_static_snapshot()
    wb = openpyxl.load_workbook(WESTERN_PADDY, data_only=True, read_only=True)
    ws = wb["WP_Paddy"]
    report = validate_sheet(parse_sheet(ws, catalog), catalog, expected_province="Western")
    assert report.would_load, [f.message for f in (report.sheet_findings + report.row_findings + report.cell_findings)]
    assert report.rows_resolved == 40
    assert report.outstanding_divisions == []


@pytest.mark.skipif(not MAX_AFFECTED.exists(), reason="Data sets/ workbook not present in this checkout")
def test_real_prose_header_sheet_is_rejected_with_named_reasons():
    catalog = load_static_snapshot()
    wb = openpyxl.load_workbook(MAX_AFFECTED, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    report = validate_sheet(parse_sheet(ws, catalog), catalog)
    assert not report.would_load
    unresolved_headers = {f.message.split("'")[1] for f in report.sheet_findings if f.code == "V011"}
    assert "Max # of drought affected People - 2005 to 2022" in unresolved_headers
