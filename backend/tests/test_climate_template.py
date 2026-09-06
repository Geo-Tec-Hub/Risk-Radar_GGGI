"""Offline tests for the CLIMATE workbook path (province-wide hazard variables).

No database: these cover the parts that decide whether a climate file is even
recognised as one, plus the generated files themselves. The authorisation rule
lives in `may_write_hazard()` in the schema and is exercised against a real
database, not here.

Run: python -m pytest backend/tests/test_climate_template.py -v
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app.importer.template_reader import read_meta, read_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
CLIMATE_DIR = REPO_ROOT / "design" / "templates" / "generated" / "_climate"
GENERATED = REPO_ROOT / "design" / "templates" / "generated"

# Provinces with no landslide profile anywhere in seed_all.sql. Their climate
# file must not ask for landslide columns -- an officer given a column with no
# profile behind it either leaves it blank forever or invents a number.
NO_LANDSLIDE = {"EAS", "NCE", "NOR"}
LANDSLIDE_COLS = {
    "LANDSLIDE_EVENTS_1974_TO_2023",
    "CUTTING_FAILURES_EARTH_SLIPS_2000_TO_2023",
    "LANDSLIDE_HAZARD_INDEX",
}


def climate_files() -> list[Path]:
    return sorted(CLIMATE_DIR.glob("CLIMATE_*_upload_template.xlsx"))


def test_all_nine_provinces_generated():
    assert len(climate_files()) == 9


@pytest.mark.parametrize("path", climate_files(), ids=lambda p: p.stem)
def test_meta_declares_climate_kind(path: Path):
    """`kind` is the whole contract. Without it the loader treats the file as a
    sector workbook, looks for a profile named by an empty sector, and refuses
    with a message about profiles that would send the reader hunting in the
    wrong place entirely."""
    meta = read_meta(openpyxl.load_workbook(path, read_only=True))
    assert meta["kind"] == "climate"
    # Excel gives an empty cell back as None, so "no sector" is falsy, not "".
    assert not meta["main_sector"]
    assert not meta["subsector"]
    assert not meta["hazard"]
    assert meta["province"]


@pytest.mark.parametrize("path", climate_files(), ids=lambda p: p.stem)
def test_reader_reports_kind(path: Path):
    wb = read_workbook(str(path))
    assert wb.errors == []
    assert wb.kind == "climate"
    assert wb.main_sector in ("", None)


@pytest.mark.parametrize("path", climate_files(), ids=lambda p: p.stem)
def test_landslide_columns_only_where_landslide_exists(path: Path):
    prov = path.stem.split("_")[1]
    cols = set(read_meta(openpyxl.load_workbook(path, read_only=True))
               ["expected_columns"].split("|"))
    if prov in NO_LANDSLIDE:
        assert not (cols & LANDSLIDE_COLS), "%s has no landslide profiles" % prov
    else:
        assert LANDSLIDE_COLS <= cols


@pytest.mark.parametrize("path", climate_files(), ids=lambda p: p.stem)
def test_no_exposure_columns(path: Path):
    """A climate file carrying a sector's own variable would write that sector's
    data under nobody's ownership. The loader refuses such a column against the
    catalogue; this catches it one step earlier, in the generator."""
    cols = read_meta(openpyxl.load_workbook(path, read_only=True))["expected_columns"].split("|")
    fixed = {"DS_CODE", "DS_DIVISION", "DISTRICT", "YEAR_START", "YEAR_END",
             "DATA_SOURCE", "NOTES"}
    variables = [c for c in cols if c not in fixed]
    assert variables
    for code in variables:
        assert any(k in code for k in
                   ("DROUGHT", "FLOOD", "LANDSLIDE", "WARM_DAYS", "PRECIPITATION",
                    "WET_DAYS", "LIGHTNING", "WIND", "CUTTING_FAILURES")), code


@pytest.mark.parametrize("path", climate_files(), ids=lambda p: p.stem)
def test_every_division_of_the_province_has_a_row(path: Path):
    meta = read_meta(openpyxl.load_workbook(path, read_only=True))
    wb = read_workbook(str(path))
    # Both period tabs, same division list on each.
    assert len(wb.sheets) == 2
    for sheet in wb.sheets:
        assert sheet.rows_scanned == int(meta["n_divisions"])


def test_division_rows_total_340_across_provinces():
    """The register has 340 DS divisions. A climate file per province must
    partition them exactly -- a division missing from every file is a division
    whose weather nobody can enter."""
    total = sum(int(read_meta(openpyxl.load_workbook(p, read_only=True))["n_divisions"])
                for p in climate_files())
    assert total == 340


def test_sector_templates_still_read_as_profile_kind():
    """`kind` defaults to "profile", so the 243 workbooks generated before
    climate files existed keep their meaning without being regenerated."""
    sample = next(GENERATED.glob("Central/*_upload_template.xlsx"))
    wb = read_workbook(str(sample))
    assert wb.kind == "profile"
    assert wb.main_sector
